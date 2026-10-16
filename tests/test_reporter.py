"""
Unit tests for src/webtest/reporter.py (3-Sheet Excel Reporter Engine).
Verifies Summary KPI formulas, TestResults 11 columns + relative hyperlinks + feedback preservation,
and NoteBug FAIL filtering + F12 Root Cause + 200x112px embedded thumbnail + 125pt row height.
"""

from pathlib import Path

from openpyxl import load_workbook
from PIL import Image as PILImage

from src.webtest.models import TestCase, TestResult
from src.webtest.reporter import (
    NOTE_BUG_HEADERS,
    NOTEBUG_ROW_HEIGHT_PT,
    TEST_RESULTS_HEADERS,
    THUMBNAIL_HEIGHT_PX,
    THUMBNAIL_WIDTH_PX,
    ExcelReporter,
)


def _create_dummy_screenshot(path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    img = PILImage.new("RGB", (800, 450), color=(31, 78, 120))
    img.save(path, format="PNG")
    return path


def test_excel_reporter_generates_3_sheets_with_all_specs(tmp_path: Path) -> None:
    screenshots_dir = tmp_path / "reports" / "screenshots"
    shot1 = _create_dummy_screenshot(screenshots_dir / "TC_01.png")
    shot2 = _create_dummy_screenshot(screenshots_dir / "TC_02.png")

    report_file = tmp_path / "reports" / "test_report_20261016.xlsx"

    tc1 = TestCase(
        id="TC_01",
        module="Authentication",
        feature="Đăng nhập",
        description="Kiểm tra đăng nhập thành công",
        steps=["Mở trang login", "Nhập thông tin hợp lệ", "Bấm Login"],
        expected_result="Vào trang Dashboard",
    )
    tc2 = TestCase(
        id="TC_02",
        module="Order",
        feature="Tạo đơn hàng",
        description="Kiểm tra tạo đơn hàng mới",
        steps=["Mở form tạo đơn", "Điền thông tin", "Bấm Lưu"],
        expected_result="Tạo đơn thành công",
    )

    res1 = TestResult(
        test_case=tc1,
        status="PASS",
        actual_result="Đăng nhập thành công vào Dashboard",
        screenshot_path=str(shot1),
        response_note="QA đã kiểm tra",
        feedback="OK",
    )
    res2 = TestResult(
        test_case=tc2,
        status="FAIL",
        actual_result="Bấm Lưu báo lỗi hệ thống 500",
        screenshot_path=str(shot2),
        ui_error="Popup báo lỗi Internal Server Error",
        console_errors=["Uncaught Error: Order submission failed at bundle.js:412"],
        network_errors=["POST /api/v1/orders -> HTTP 500 Internal Server Error"],
    )

    reporter = ExcelReporter(
        report_path=report_file,
        suite_name="Test Suite Chuẩn Doanh Nghiệp",
        target_url="https://app.example.com",
    )
    reporter.add_results([res1, res2])
    saved_path = reporter.save()

    assert saved_path.exists()

    wb = load_workbook(saved_path, data_only=False)
    assert wb.sheetnames == ["Summary", "TestResults", "NoteBug"]

    # 1. Verify Sheet 'Summary'
    ws_sum = wb["Summary"]
    assert ws_sum["A1"].value == "BÁO CÁO TỔNG QUAN KIỂM THỬ (AUTOMATION TEST SUMMARY)"
    assert ws_sum["A1"].fill.start_color.rgb.endswith("1F4E78")
    assert ws_sum["C7"].value == "=COUNTA(TestResults!A2:A100)"
    assert ws_sum["C8"].value == '=COUNTIF(TestResults!I2:I100, "PASS")'
    assert ws_sum["C9"].value == '=COUNTIF(TestResults!I2:I100, "FAIL")'
    assert ws_sum["C10"].value == "=C8/C7"
    assert ws_sum["C10"].number_format == "0.0%"
    assert ws_sum["C11"].value == "FAILED"

    # 2. Verify Sheet 'TestResults'
    ws_tr = wb["TestResults"]
    assert ws_tr.freeze_panes == "A2"
    assert ws_tr.auto_filter.ref == "A1:K3"
    actual_tr_headers = [ws_tr.cell(row=1, column=i).value for i in range(1, 12)]
    assert actual_tr_headers == TEST_RESULTS_HEADERS

    # No floating images allowed in TestResults
    assert len(ws_tr._images) == 0

    # Check relative hyperlink in Column 8 (H2 and H3)
    cell_h2 = ws_tr["H2"]
    assert cell_h2.value == "Xem ảnh: TC_01.png"
    assert cell_h2.hyperlink is not None
    assert cell_h2.hyperlink.target == "reports/screenshots/TC_01.png"

    # Check PASS / FAIL colors in Column 9 (I2 and I3)
    assert ws_tr["I2"].value == "PASS"
    assert ws_tr["I2"].fill.start_color.rgb.endswith("C6EFCE")
    assert ws_tr["I2"].font.color.rgb.endswith("006100")

    assert ws_tr["I3"].value == "FAIL"
    assert ws_tr["I3"].fill.start_color.rgb.endswith("FFC7CE")
    assert ws_tr["I3"].font.color.rgb.endswith("9C0006")

    # 3. Verify Sheet 'NoteBug'
    ws_bug = wb["NoteBug"]
    assert ws_bug.freeze_panes == "A2"
    actual_bug_headers = [ws_bug.cell(row=1, column=i).value for i in range(1, 11)]
    assert actual_bug_headers == NOTE_BUG_HEADERS

    # Only 1 FAIL case (TC_02) should be in NoteBug (row 2)
    assert ws_bug["A2"].value == "TC_02"
    assert ws_bug["A3"].value is None
    assert "Uncaught Error: Order submission failed" in ws_bug["G2"].value
    assert "POST /api/v1/orders -> HTTP 500" in ws_bug["G2"].value
    assert ws_bug["I2"].value == "Chờ xác nhận"

    # Check row height is 125pt and image is embedded with 200x112px
    assert ws_bug.row_dimensions[2].height == NOTEBUG_ROW_HEIGHT_PT
    assert len(ws_bug._images) == 1
    embedded_img = ws_bug._images[0]
    assert embedded_img.width == THUMBNAIL_WIDTH_PX
    assert embedded_img.height == THUMBNAIL_HEIGHT_PX

    wb.close()


def test_excel_reporter_preserves_user_feedback_on_incremental_update(tmp_path: Path) -> None:
    report_file = tmp_path / "reports" / "incremental_report.xlsx"

    tc1 = TestCase(
        id="TC_01",
        module="Auth",
        feature="Login",
        description="Kiểm tra login",
        steps=["Mở trang", "Đăng nhập"],
        expected_result="Thành công",
    )

    # First run: log TC_01 with user feedback
    reporter1 = ExcelReporter(report_path=report_file)
    reporter1.log_result(
        TestResult(
            test_case=tc1,
            status="FAIL",
            actual_result="Lỗi lần 1",
            response_note="Dev đang xử lý",
            feedback="Ưu tiên cao",
            bug_status="Đang fix",
        )
    )

    # Second run: re-open existing workbook and update TC_01 without specifying feedback
    reporter2 = ExcelReporter(report_path=report_file)
    reporter2.log_result(
        TestResult(
            test_case=tc1,
            status="PASS",
            actual_result="Đã chạy lại thành công sau khi hotfix",
        )
    )

    wb = load_workbook(report_file, data_only=False)
    ws_tr = wb["TestResults"]
    assert ws_tr["I2"].value == "PASS"
    assert ws_tr["J2"].value == "Dev đang xử lý"
    assert ws_tr["K2"].value == "Ưu tiên cao"
    assert wb["Summary"]["C11"].value == "PASSED"
    wb.close()
