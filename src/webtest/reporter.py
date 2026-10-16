"""
Enterprise 3-Sheet Excel Reporter Engine for AI Web Testing Hybrid Architecture.

Generates and updates an Excel workbook (.xlsx) with 3 standardized sheets:
1. Summary: KPI Dashboard with dynamic Excel formulas and Module breakdown.
2. TestResults: 11-column company standard sheet with Freeze Panes A2, Auto-Filter,
   relative screenshot hyperlinks (NO floating images), and preserved user feedback columns.
3. NoteBug: Dedicated FAIL bug tracker (10 columns) with F12 Console & Network API Root Cause,
   embedded 200x112px screenshot thumbnails in Column H, and 125pt row height.
"""

from __future__ import annotations

from collections import OrderedDict
from datetime import datetime
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .models import ErrorCode, TestCase, TestResult, TestStatus

# Color Palette Constants (Enterprise QA Standard)
NAVY_HEADER_HEX = "1F4E78"
SUBHEADER_HEX = "D9E1F2"
BUG_HEADER_HEX = "8B0000"
PASS_BG_HEX = "C6EFCE"
PASS_FG_HEX = "006100"
FAIL_BG_HEX = "FFC7CE"
FAIL_FG_HEX = "9C0006"
ZEBRA_BG_HEX = "F9FAFB"
WHITE_HEX = "FFFFFF"

TEST_RESULTS_HEADERS: list[str] = [
    "Test Case ID",
    "Module",
    "Chức năng",
    "Mô tả",
    "Các bước thực hiện",
    "Expected Result",
    "Actual Result / Đánh giá",
    "Log/Screenshot/Video",
    "Status",
    "Phản hồi",
    "FEEDBACK",
]

NOTE_BUG_HEADERS: list[str] = [
    "Bug ID / Test Case ID",
    "Module",
    "Chức năng",
    "Các bước tái hiện",
    "Expected Result",
    "Actual Result",
    "Chi tiết lỗi (Root Cause)",
    "Bằng chứng ảnh (Screenshot)",
    "Trạng thái xử lý",
    "Dev Note / Phản hồi",
]

THUMBNAIL_WIDTH_PX = 200
THUMBNAIL_HEIGHT_PX = 112
NOTEBUG_ROW_HEIGHT_PT = 125


def _thin_border() -> Border:
    side = Side(border_style="thin", color="D9D9D9")
    return Border(left=side, right=side, top=side, bottom=side)


class ExcelReporter:
    """
    Manages creation and incremental updates of the 3-sheet QA Automation Excel Report.
    """

    def __init__(
        self,
        report_path: str | Path,
        suite_name: str = "AI Web Testing Suite (Chrome DevTools MCP)",
        target_url: str = "",
    ) -> None:
        self.report_path = Path(report_path)
        self.suite_name = suite_name
        self.target_url = target_url
        self._results: OrderedDict[str, TestResult] = OrderedDict()

        if self.report_path.exists():
            self._load_existing_workbook()

    @property
    def results(self) -> list[TestResult]:
        """Return all currently tracked TestResult objects in insertion order."""
        return list(self._results.values())

    def _resolve_screenshot_file(self, screenshot_path: str) -> Path | None:
        """Resolve a screenshot path relative to CWD, report directory, or project root."""
        if not screenshot_path:
            return None
        raw_path = Path(screenshot_path)
        candidates = [
            raw_path,
            self.report_path.parent / raw_path,
            self.report_path.parent / "screenshots" / raw_path.name,
            self.report_path.parent.parent / raw_path,
        ]
        for candidate in candidates:
            if candidate.exists() and candidate.is_file():
                return candidate.resolve()
        return None

    @staticmethod
    def _normalize_relative_screenshot_link(screenshot_path: str) -> str:
        """
        Normalize the hyperlink target for Column 8 of TestResults to relative format:
        'reports/screenshots/<filename>' (or the normalized relative path).
        """
        cleaned = screenshot_path.strip().replace("\\", "/")
        if not cleaned:
            return ""
        if "reports/screenshots/" in cleaned:
            idx = cleaned.find("reports/screenshots/")
            return cleaned[idx:]
        filename = Path(cleaned).name
        return f"reports/screenshots/{filename}"

    def _load_existing_workbook(self) -> None:
        """
        Read existing rows from 'TestResults' and 'NoteBug' so incremental CLI calls
        preserve prior test cases as well as user-filled 'Phản hồi' and 'FEEDBACK'.
        """
        try:
            wb = load_workbook(self.report_path, data_only=False)
        except Exception:
            return

        bug_lookup: dict[str, tuple[str, str, str]] = {}
        if "NoteBug" in wb.sheetnames:
            ws_bug = wb["NoteBug"]
            for row_idx in range(2, ws_bug.max_row + 1):
                bug_id = str(ws_bug.cell(row=row_idx, column=1).value or "").strip()
                if not bug_id:
                    continue
                root_cause_val = str(ws_bug.cell(row=row_idx, column=7).value or "").strip()
                bug_status_val = str(ws_bug.cell(row=row_idx, column=9).value or "Chờ xác nhận").strip()
                dev_note_val = str(ws_bug.cell(row=row_idx, column=10).value or "").strip()
                bug_lookup[bug_id] = (root_cause_val, bug_status_val, dev_note_val)

        if "TestResults" in wb.sheetnames:
            ws_tr = wb["TestResults"]
            for row_idx in range(2, ws_tr.max_row + 1):
                tc_id = str(ws_tr.cell(row=row_idx, column=1).value or "").strip()
                if not tc_id:
                    continue
                module = str(ws_tr.cell(row=row_idx, column=2).value or "General").strip()
                feature = str(ws_tr.cell(row=row_idx, column=3).value or tc_id).strip()
                description = str(ws_tr.cell(row=row_idx, column=4).value or feature).strip()
                steps_raw = str(ws_tr.cell(row=row_idx, column=5).value or "1. Thực hiện kiểm thử").strip()
                steps = [line.strip() for line in steps_raw.splitlines() if line.strip()] or ["1. Thực hiện kiểm thử"]
                expected = str(ws_tr.cell(row=row_idx, column=6).value or "Hoạt động đúng").strip()
                actual = str(ws_tr.cell(row=row_idx, column=7).value or "Đã thực thi").strip()

                link_cell = ws_tr.cell(row=row_idx, column=8)
                screenshot_path = ""
                if link_cell.hyperlink and link_cell.hyperlink.target:
                    screenshot_path = str(link_cell.hyperlink.target).strip()
                elif link_cell.value and "Xem ảnh:" in str(link_cell.value):
                    fname = str(link_cell.value).split("Xem ảnh:", 1)[1].strip()
                    screenshot_path = f"reports/screenshots/{fname}"

                status_val = str(ws_tr.cell(row=row_idx, column=9).value or "PASS").strip().upper()
                if status_val not in ("PASS", "FAIL"):
                    status_val = "PASS"

                response_note = str(ws_tr.cell(row=row_idx, column=10).value or "").strip()
                feedback = str(ws_tr.cell(row=row_idx, column=11).value or "").strip()

                root_cause_val, bug_status_val, dev_note_val = bug_lookup.get(
                    tc_id, ("", "Chờ xác nhận", "")
                )

                tc = TestCase(
                    id=tc_id,
                    module=module,
                    feature=feature,
                    description=description,
                    steps=steps,
                    expected_result=expected,
                )
                self._results[tc_id] = TestResult(
                    test_case=tc,
                    status=status_val,
                    actual_result=actual,
                    screenshot_path=screenshot_path,
                    root_cause=root_cause_val,
                    bug_status=bug_status_val,
                    response_note=response_note or dev_note_val,
                    feedback=feedback,
                )

        wb.close()

    def add_result(self, result: TestResult) -> None:
        """
        Add or update a TestResult in memory while preserving existing user-entered
        'Phản hồi' (response_note), 'FEEDBACK' (feedback), and 'Trạng thái xử lý' (bug_status).
        """
        tc_id = result.test_case.id
        if tc_id in self._results:
            existing = self._results[tc_id]
            if not result.response_note and existing.response_note:
                result.response_note = existing.response_note
            if not result.feedback and existing.feedback:
                result.feedback = existing.feedback
            if (
                result.bug_status == "Chờ xác nhận"
                and existing.bug_status
                and existing.bug_status != "Chờ xác nhận"
            ):
                result.bug_status = existing.bug_status

        self._results[tc_id] = result

    def add_results(self, results: Iterable[TestResult]) -> None:
        """Add or update multiple TestResult objects."""
        for res in results:
            self.add_result(res)

    def log_result(self, result: TestResult) -> Path:
        """Add/update a single TestResult and immediately save the 3-sheet workbook."""
        self.add_result(result)
        return self.save()

    def save(self) -> Path:
        """Generate and save the complete 3-sheet Excel report to disk."""
        self.report_path.parent.mkdir(parents=True, exist_ok=True)

        wb = Workbook()
        # Sheet 1: Summary
        ws_summary = wb.active
        ws_summary.title = "Summary"
        # Sheet 2: TestResults
        ws_results = wb.create_sheet(title="TestResults")
        # Sheet 3: NoteBug
        ws_notebug = wb.create_sheet(title="NoteBug")

        self._build_summary_sheet(ws_summary)
        self._build_test_results_sheet(ws_results)
        self._build_notebug_sheet(ws_notebug)

        wb.save(self.report_path)
        wb.close()
        return self.report_path

    def _build_summary_sheet(self, ws: Worksheet) -> None:
        """
        Build Sheet 1: Summary (KPI Dashboard)
        - Banner in A1:E2 with Navy background (1F4E78) and white bold text.
        - KPI block at rows 6..11 so that:
          - C7 = '=COUNTA(TestResults!A2:A100)' (Tổng số Test Case)
          - C8 = '=COUNTIF(TestResults!I2:I100, "PASS")' (Số lượng PASS)
          - C9 = '=COUNTIF(TestResults!I2:I100, "FAIL")' (Số lượng FAIL)
          - C10 = '=C8/C7' formatted as '0.0%' (Tỷ lệ Pass Rate %)
          - C11 = 'PASSED' if 100% Pass else 'FAILED' (Đánh giá chung)
        - Module breakdown table starting at row 13.
        """
        border = _thin_border()
        navy_fill = PatternFill(start_color=NAVY_HEADER_HEX, end_color=NAVY_HEADER_HEX, fill_type="solid")
        sub_fill = PatternFill(start_color=SUBHEADER_HEX, end_color=SUBHEADER_HEX, fill_type="solid")
        pass_fill = PatternFill(start_color=PASS_BG_HEX, end_color=PASS_BG_HEX, fill_type="solid")
        fail_fill = PatternFill(start_color=FAIL_BG_HEX, end_color=FAIL_BG_HEX, fill_type="solid")

        # 1. Banner Title (A1:E2)
        ws.merge_cells("A1:E2")
        banner_cell = ws["A1"]
        banner_cell.value = "BÁO CÁO TỔNG QUAN KIỂM THỬ (AUTOMATION TEST SUMMARY)"
        banner_cell.fill = navy_fill
        banner_cell.font = Font(name="Calibri", size=14, bold=True, color=WHITE_HEX)
        banner_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 24
        ws.row_dimensions[2].height = 24

        # 2. Metadata rows (Rows 3-5)
        ws["B3"] = "Bộ kịch bản:"
        ws["C3"] = self.suite_name
        ws["B4"] = "URL Kiểm thử:"
        ws["C4"] = self.target_url or "N/A"
        ws["B5"] = "Thời gian xuất báo cáo:"
        ws["C5"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        for r in range(3, 6):
            ws.cell(row=r, column=2).font = Font(bold=True)

        # 3. KPI Block Header (Row 6)
        ws["B6"] = "CHỈ SỐ ĐO LƯỜNG (KPI)"
        ws["C6"] = "GIÁ TRỊ"
        for col_idx in (2, 3):
            cell = ws.cell(row=6, column=col_idx)
            cell.fill = navy_fill
            cell.font = Font(bold=True, color=WHITE_HEX)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border

        all_results = self.results
        total_count = len(all_results)
        fail_count = sum(1 for r in all_results if r.is_failed)
        overall_evaluation = "PASSED" if (total_count > 0 and fail_count == 0) else "FAILED"

        kpi_rows = [
            (7, "Tổng số Test Case", "=COUNTA(TestResults!A2:A100)", None),
            (8, "Số lượng PASS", '=COUNTIF(TestResults!I2:I100, "PASS")', None),
            (9, "Số lượng FAIL", '=COUNTIF(TestResults!I2:I100, "FAIL")', None),
            (10, "Tỷ lệ Pass Rate %", "=C8/C7", "0.0%"),
            (11, "Đánh giá chung", overall_evaluation, None),
        ]

        for row_idx, label, formula_or_val, num_format in kpi_rows:
            lbl_cell = ws.cell(row=row_idx, column=2, value=label)
            val_cell = ws.cell(row=row_idx, column=3, value=formula_or_val)
            lbl_cell.font = Font(bold=True)
            lbl_cell.border = border
            val_cell.font = Font(bold=True)
            val_cell.alignment = Alignment(horizontal="center", vertical="center")
            val_cell.border = border

            if num_format:
                val_cell.number_format = num_format

            if row_idx == 8:
                val_cell.fill = pass_fill
                val_cell.font = Font(bold=True, color=PASS_FG_HEX)
            elif row_idx == 9:
                val_cell.fill = fail_fill
                val_cell.font = Font(bold=True, color=FAIL_FG_HEX)
            elif row_idx == 11:
                if overall_evaluation == "PASSED":
                    val_cell.fill = pass_fill
                    val_cell.font = Font(bold=True, color=PASS_FG_HEX)
                else:
                    val_cell.fill = fail_fill
                    val_cell.font = Font(bold=True, color=FAIL_FG_HEX)

        # 4. Module Breakdown Table (Starting at Row 13)
        ws["A13"] = "THỐNG KÊ THEO PHÂN HỆ (MODULE BREAKDOWN)"
        ws.merge_cells("A13:E13")
        ws["A13"].fill = sub_fill
        ws["A13"].font = Font(bold=True, size=11, color=NAVY_HEADER_HEX)
        ws["A13"].alignment = Alignment(horizontal="left", vertical="center")

        module_headers = ["Module", "Tổng số Case", "Số lượng PASS", "Số lượng FAIL", "Pass Rate (%)"]
        for col_idx, header_text in enumerate(module_headers, start=1):
            c = ws.cell(row=14, column=col_idx, value=header_text)
            c.fill = navy_fill
            c.font = Font(bold=True, color=WHITE_HEX)
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = border

        module_stats: OrderedDict[str, dict[str, int]] = OrderedDict()
        for res in all_results:
            mod = res.test_case.module
            if mod not in module_stats:
                module_stats[mod] = {"total": 0, "pass": 0, "fail": 0}
            module_stats[mod]["total"] += 1
            if res.is_passed:
                module_stats[mod]["pass"] += 1
            else:
                module_stats[mod]["fail"] += 1

        current_row = 15
        for mod_name, stats in module_stats.items():
            tot = stats["total"]
            pas = stats["pass"]
            fai = stats["fail"]
            rate = (pas / tot) if tot > 0 else 0.0

            row_vals = [mod_name, tot, pas, fai, rate]
            for col_idx, val in enumerate(row_vals, start=1):
                c = ws.cell(row=current_row, column=col_idx, value=val)
                c.border = border
                c.alignment = Alignment(
                    horizontal="left" if col_idx == 1 else "center",
                    vertical="center",
                )
                if col_idx == 5:
                    c.number_format = "0.0%"
            current_row += 1

        col_widths = {1: 24, 2: 28, 3: 28, 4: 18, 5: 18}
        for col_idx, width in col_widths.items():
            ws.column_dimensions[get_column_letter(col_idx)].width = width

    def _build_test_results_sheet(self, ws: Worksheet) -> None:
        """
        Build Sheet 2: TestResults (11-column company standard)
        - Freeze Panes at A2, Auto-Filter enabled.
        - Column 8 contains relative hyperlink ONLY (NO floating images).
        - Column 9 colored by Status (PASS/FAIL).
        - Columns 10 (Phản hồi) & 11 (FEEDBACK) preserved.
        """
        border = _thin_border()
        header_fill = PatternFill(start_color=NAVY_HEADER_HEX, end_color=NAVY_HEADER_HEX, fill_type="solid")
        pass_fill = PatternFill(start_color=PASS_BG_HEX, end_color=PASS_BG_HEX, fill_type="solid")
        fail_fill = PatternFill(start_color=FAIL_BG_HEX, end_color=FAIL_BG_HEX, fill_type="solid")
        zebra_fill = PatternFill(start_color=ZEBRA_BG_HEX, end_color=ZEBRA_BG_HEX, fill_type="solid")

        # Row 1: 11 Bold Headers
        ws.row_dimensions[1].height = 28
        for col_idx, title in enumerate(TEST_RESULTS_HEADERS, start=1):
            cell = ws.cell(row=1, column=col_idx, value=title)
            cell.fill = header_fill
            cell.font = Font(name="Calibri", size=11, bold=True, color=WHITE_HEX)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border

        ws.freeze_panes = "A2"

        all_results = self.results
        for idx, res in enumerate(all_results, start=2):
            tc = res.test_case
            row_values = [
                tc.id,
                tc.module,
                tc.feature,
                tc.description,
                tc.formatted_steps(),
                tc.expected_result,
                res.actual_result,
                "",  # Column 8 populated with relative hyperlink below
                res.status,
                res.response_note,
                res.feedback,
            ]

            for col_idx, val in enumerate(row_values, start=1):
                cell = ws.cell(row=idx, column=col_idx, value=val)
                cell.border = border
                cell.alignment = Alignment(
                    horizontal="center" if col_idx in (1, 8, 9) else "left",
                    vertical="top",
                    wrap_text=True,
                )
                if idx % 2 == 1 and col_idx != 9:
                    cell.fill = zebra_fill

            # Column 8: Relative Hyperlink ONLY (strictly no embedded image in TestResults)
            link_cell = ws.cell(row=idx, column=8)
            if res.screenshot_path:
                rel_target = self._normalize_relative_screenshot_link(res.screenshot_path)
                filename = Path(rel_target).name
                link_cell.value = f"Xem ảnh: {filename}"
                link_cell.hyperlink = rel_target
                link_cell.font = Font(color="0563C1", underline="single")
            else:
                link_cell.value = "N/A"

            # Column 9: Status Styling
            status_cell = ws.cell(row=idx, column=9)
            if res.is_passed:
                status_cell.fill = pass_fill
                status_cell.font = Font(bold=True, color=PASS_FG_HEX)
            else:
                status_cell.fill = fail_fill
                status_cell.font = Font(bold=True, color=FAIL_FG_HEX)

            ws.row_dimensions[idx].height = 48

        last_row = max(len(all_results) + 1, 2)
        ws.auto_filter.ref = f"A1:K{last_row}"

        col_widths = {
            1: 16,
            2: 20,
            3: 24,
            4: 30,
            5: 38,
            6: 32,
            7: 32,
            8: 26,
            9: 14,
            10: 24,
            11: 24,
        }
        for col_idx, width in col_widths.items():
            ws.column_dimensions[get_column_letter(col_idx)].width = width

    def _build_notebug_sheet(self, ws: Worksheet) -> None:
        """
        Build Sheet 3: NoteBug (FAIL cases only, 10 columns)
        - Column 7: Detailed Root Cause (UI + Console F12 + Network API).
        - Column 8 (Column H): Directly embedded thumbnail image (200x112px) using openpyxl.drawing.image.Image.
        - Row height automatically set to 125pt for every bug row.
        """
        border = _thin_border()
        header_fill = PatternFill(start_color=BUG_HEADER_HEX, end_color=BUG_HEADER_HEX, fill_type="solid")

        ws.row_dimensions[1].height = 28
        for col_idx, title in enumerate(NOTE_BUG_HEADERS, start=1):
            cell = ws.cell(row=1, column=col_idx, value=title)
            cell.fill = header_fill
            cell.font = Font(name="Calibri", size=11, bold=True, color=WHITE_HEX)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border

        ws.freeze_panes = "A2"

        failed_results = [r for r in self.results if r.is_failed]
        for idx, res in enumerate(failed_results, start=2):
            tc = res.test_case
            root_cause_text = res.build_root_cause_details()
            row_values = [
                tc.id,
                tc.module,
                tc.feature,
                tc.formatted_steps(),
                tc.expected_result,
                res.actual_result,
                root_cause_text,
                "",  # Column 8 (H): Embedded thumbnail image
                res.bug_status or "Chờ xác nhận",
                res.response_note,
            ]

            for col_idx, val in enumerate(row_values, start=1):
                cell = ws.cell(row=idx, column=col_idx, value=val)
                cell.border = border
                cell.alignment = Alignment(
                    horizontal="center" if col_idx in (1, 8, 9) else "left",
                    vertical="top",
                    wrap_text=True,
                )

            # Set mandatory 125pt row height for NoteBug rows
            ws.row_dimensions[idx].height = NOTEBUG_ROW_HEIGHT_PT

            # Embed 200x112px thumbnail image into cell H{idx} if screenshot file exists on disk
            img_file = self._resolve_screenshot_file(res.screenshot_path)
            if img_file is not None:
                excel_img = OpenpyxlImage(str(img_file))
                excel_img.width = THUMBNAIL_WIDTH_PX
                excel_img.height = THUMBNAIL_HEIGHT_PX
                ws.add_image(excel_img, f"H{idx}")
            else:
                ws.cell(row=idx, column=8, value=res.screenshot_path or "Chưa có ảnh")

        last_row = max(len(failed_results) + 1, 2)
        ws.auto_filter.ref = f"A1:J{last_row}"

        col_widths = {
            1: 20,
            2: 20,
            3: 24,
            4: 36,
            5: 30,
            6: 30,
            7: 48,
            8: 30,
            9: 18,
            10: 26,
        }
        for col_idx, width in col_widths.items():
            ws.column_dimensions[get_column_letter(col_idx)].width = width
